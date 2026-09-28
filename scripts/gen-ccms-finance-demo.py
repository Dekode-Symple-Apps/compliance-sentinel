"""The Finance review demo: a supply agreement over RM500,000 (so not
straight-through) with a full payment schedule — advance against a bond,
milestone payments, retention — for Finance to validate and clear while the
contract is in review.

  python3 scripts/gen-ccms-finance-demo.py →
    ~/Desktop/01. Demo Data/Commercial CMS/13 Supply Agreement - Sinar Precast - Block D (payment terms).docx
Invented details; entity particulars are the demo values in src/lib/ccms.ts.
"""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.expanduser("~/Desktop/01. Demo Data/Commercial CMS/13 Supply Agreement - Sinar Precast - Block D (payment terms).docx")

CLAUSES = [
    ("Supply", [
        "The Supplier shall supply the precast concrete beams, hollow-core slabs and staircase units described in Schedule 1 (the \"Goods\") to the LSH 33 Block D site, Jalan Klang Lama, Kuala Lumpur.",
        "The Goods shall comply with MS 1314 and the approved shop drawings, be SIRIM certified and carry a valid CIDB Perakuan Pematuhan Standard (PPS). Each delivery shall be accompanied by test certificates.",
    ]),
    ("Price", ["The contract price is RM1,280,000 (Ringgit Malaysia One Million Two Hundred and Eighty Thousand), fixed and exclusive of SST, including delivery, unloading and installation supervision."]),
    ("Delivery", [
        "The Supplier shall deliver the Goods in three batches: Batch 1 by 15 December 2026, Batch 2 by 29 January 2027 and Batch 3 by 12 March 2027.",
        "The Purchaser may reschedule a batch on 14 days' written notice without additional cost if the new date is within 60 days of the original date.",
        "If a batch is late, the Supplier shall pay liquidated damages of 0.1% of the contract price per day of delay, capped at 10% of the contract price.",
    ]),
    ("Payment", [
        "The Purchaser shall pay the contract price in the instalments in Schedule 2, each within 30 days of receiving a valid invoice issued after the milestone is met.",
        "The advance payment is due only after the Supplier delivers an advance payment bond for the same amount, from a licensed Malaysian bank, valid until Batch 1 is accepted.",
        "The Purchaser shall retain 5% of the contract price as retention, released 12 months after acceptance of the last batch, less any amounts owed by the Supplier.",
        "The Purchaser may set off against any payment any liquidated damages or other amounts the Supplier owes under this Agreement.",
    ]),
    ("Performance Bond", ["Within 14 days of signing, the Supplier shall deliver a performance bond of 5% of the contract price (RM64,000) from a licensed Malaysian bank, valid until 3 months after the end of the warranty period."]),
    ("Inspection, Title and Warranty", [
        "The Purchaser may reject non-conforming Goods by written notice within 7 days of delivery; the Supplier shall replace them within 14 days at its own cost, failing which the Purchaser may buy substitutes and recover the extra cost.",
        "Title and risk pass to the Purchaser on acceptance at site.",
        "The Supplier warrants the Goods for 24 months from acceptance and shall remedy defects within 14 days of notice.",
    ]),
    ("Insurance", ["The Supplier shall maintain goods-in-transit and product liability insurance of at least RM2,000,000 per occurrence, and provide certificates before the first delivery."]),
    ("Indemnity and Liability", [
        "The Supplier shall indemnify the Purchaser against third-party claims arising from defects in the Goods or the Supplier's negligence.",
        "Each party's total liability is limited to the contract price, except for death or personal injury, damage to property, fraud, gross negligence, wilful misconduct, the indemnity above, or breach of the Anti-Bribery clause.",
    ]),
    ("Anti-Bribery and Corruption", ["The Supplier shall comply with the Malaysian Anti-Corruption Commission Act 2009 (including section 17A) and the Purchaser's Anti-Bribery and Corruption Policy. The Purchaser may terminate immediately for a breach of this clause."]),
    ("Confidentiality and Personal Data", ["Each party shall keep the other's confidential information confidential and comply with the Personal Data Protection Act 2010."]),
    ("Termination", [
        "Either party may terminate for a material breach not remedied within 14 days of written notice.",
        "The Purchaser may terminate immediately if the Supplier becomes insolvent or enters administration, liquidation or winding-up, and may terminate for convenience on 30 days' notice, paying for Goods accepted and Goods already manufactured for this Agreement.",
    ]),
    ("Force Majeure", ["Neither party is liable for delay caused by events beyond its reasonable control, if it notifies the other promptly and mitigates. Either party may terminate if the event continues for more than 60 days."]),
    ("Assignment and Subcontracting", ["The Supplier shall not assign this Agreement or subcontract manufacture without the Purchaser's prior written consent."]),
    ("Notices", ["Notices shall be in writing to: the Purchaser — attention Head of Procurement, Wisma Lim Seong Hai, Kuala Lumpur, procurement@lsh-best.example; the Supplier — attention Rahman bin Ali, General Manager, Lot 45, Kawasan Perindustrian Senai, Johor, rahman@sinarprecast.example."]),
    ("Stamp Duty", ["The Supplier shall bear the stamp duty on this Agreement and arrange stamping within 30 days of execution."]),
    ("Governing Law and Disputes", ["This Agreement is governed by the laws of Malaysia and the parties submit to the exclusive jurisdiction of the courts of Malaysia."]),
    ("Entire Agreement", ["This Agreement, with its Schedules, is the entire agreement between the parties and may be varied only in writing signed by both parties."]),
]

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
def h(t, size=12, center=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(size)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
para = lambda t: doc.add_paragraph(t)
h("SUPPLY AGREEMENT", 16, True)
h("Precast Structural Components — LSH 33 Block D", 12, True)
para("THIS AGREEMENT is made on 29 September 2026 and takes effect on 15 October 2026 (the \"Commencement Date\").")
h("BETWEEN")
para("(1) LSH BEST BUILDERS SDN BHD (Registration No. 200501023456 (234567-A)), of Wisma Lim Seong Hai, Kuala Lumpur (the \"Purchaser\"); and")
para("(2) SINAR PRECAST INDUSTRIES SDN BHD (Registration No. 201501012345), of Lot 45, Kawasan Perindustrian Senai, Johor (the \"Supplier\").")
for i, (title, lines) in enumerate(CLAUSES, 1):
    h(f"{i}. {title}")
    for j, l in enumerate(lines, 1): para(f"{i}.{j} {l}")
h("SCHEDULE 1 — GOODS")
t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
for i, x in enumerate(["Item", "Quantity", "Unit price (RM)", "Amount (RM)"]): t.rows[0].cells[i].text = x
for row in [("Precast beams, 6 m", "180", "3,200", "576,000"), ("Hollow-core slabs, 1.2 m x 8 m", "240", "2,400", "576,000"), ("Precast staircase units", "32", "4,000", "128,000")]:
    c = t.add_row().cells
    for i, x in enumerate(row): c[i].text = x
c = t.add_row().cells; c[0].text = "Total"; c[3].text = "1,280,000"
h("SCHEDULE 2 — PAYMENT SCHEDULE")
t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
for i, x in enumerate(["Milestone", "Percentage", "Amount (RM)", "Due"]): t.rows[0].cells[i].text = x
for row in [("Advance payment, against an advance payment bond", "10%", "128,000", "Within 30 days of the invoice issued on signing (by 29 October 2026)"),
            ("Batch 1 delivered and accepted", "30%", "384,000", "Batch 1 due 15 December 2026; payment by 14 January 2027"),
            ("Batch 2 delivered and accepted", "30%", "384,000", "Batch 2 due 29 January 2027; payment by 28 February 2027"),
            ("Batch 3 delivered and accepted", "25%", "320,000", "Batch 3 due 12 March 2027; payment by 11 April 2027"),
            ("Retention released", "5%", "64,000", "12 months after acceptance of Batch 3 (by 12 March 2028)")]:
    c = t.add_row().cells
    for i, x in enumerate(row): c[i].text = x
c = t.add_row().cells; c[0].text = "Total"; c[1].text = "100%"; c[2].text = "1,280,000"
h("EXECUTION")
para("Executed in accordance with section 66 of the Companies Act 2016.")
para("Signed for and on behalf of LSH BEST BUILDERS SDN BHD:    ____________________ Director    ____________________ Director")
para("Signed for and on behalf of SINAR PRECAST INDUSTRIES SDN BHD:    ____________________ Director    ____________________ Company Secretary")
doc.save(OUT); print(OUT)
