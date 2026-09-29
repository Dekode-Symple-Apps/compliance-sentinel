"""Signed and stamped IT service agreement with a 50/30/20 payment schedule — for
the obligations demo (File to Repository extracts one Finance obligation per
milestone). Invented parties' details; entity particulars are the demo values
in src/lib/ccms.ts.

  python3 scripts/gen-ccms-service-demo.py →
    ~/Desktop/01. Demo Data/Commercial CMS/6 Obligations - Awan Digital IT service/
      Already filed as CC-2026-0009 - signed IT service agreement (reference).docx
"""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.expanduser("~/Desktop/01. Demo Data/Commercial CMS/6 Obligations - Awan Digital IT service/Already filed as CC-2026-0009 - signed IT service agreement (reference).docx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)

def h(t, size=12, center=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(size)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
def para(t): doc.add_paragraph(t)

h("IT SERVICE AGREEMENT", 16, True)
para("Stamp duty paid — Certificate No. STAMPS-2026-DEMO-0212 · LHDN · 30 September 2026")
para("THIS AGREEMENT is made on 29 September 2026")
h("BETWEEN")
para("(1) LIM SEONG HAI LIGHTING SDN BHD (Registration No. 201101045678 (456789-C)), of Wisma Lim Seong Hai, Kuala Lumpur (the \"Company\"); and")
para("(2) AWAN DIGITAL SOLUTIONS SDN BHD (Registration No. 202001056789), of Level 8, Menara Awan, Petaling Jaya, Selangor (the \"Service Provider\").")
h("1. Services")
para("1.1 The Service Provider shall implement and support a cloud ERP system for the Company as described in Schedule 1 (the \"Services\").")
para("1.2 The Service Provider shall deliver the user acceptance testing (UAT) environment by 15 December 2026.")
para("1.3 The Service Provider shall submit a written progress report to the Company on the last business day of each month.")
h("2. Term")
para("2.1 This Agreement commences on 1 October 2026 and continues until 30 September 2027, unless terminated earlier under clause 9.")
h("3. Contract Price and Payment")
para("3.1 The contract price is RM480,000 (Ringgit Malaysia Four Hundred and Eighty Thousand), exclusive of SST.")
para("3.2 The Company shall pay the contract price in the instalments set out in Schedule 2, within 30 days of receiving a valid invoice for each instalment.")
h("4. Service Levels")
para("4.1 The Service Provider shall maintain system availability of at least 99.5% each month after go-live.")
h("5. Confidentiality and Data Protection")
para("5.1 Each party shall keep the other's confidential information confidential and comply with the Personal Data Protection Act 2010.")
para("5.2 The Service Provider shall notify the Company in writing within 72 hours of becoming aware of any personal data breach.")
h("6. Anti-Bribery")
para("6.1 Each party shall comply with the Malaysian Anti-Corruption Commission Act 2009 and the Company's Anti-Bribery and Corruption Policy.")
h("7. Liability")
para("7.1 Neither party limits its liability for fraud, wilful misconduct or breach of clauses 5 and 6.")
h("8. Renewal")
para("8.1 The parties may renew this Agreement for a further 12 months by written agreement made at least 30 days before expiry.")
h("9. Termination")
para("9.1 Either party may terminate this Agreement by 30 days' written notice.")
h("10. Governing Law")
para("10.1 This Agreement is governed by the laws of Malaysia.")
h("SCHEDULE 2 — PAYMENT SCHEDULE")
t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
for i, x in enumerate(["Milestone", "Percentage", "Amount (RM)", "Due"]): t.rows[0].cells[i].text = x
for row in [("Contract signing", "50%", "240,000", "Within 30 days of the invoice issued on signing (by 29 October 2026)"),
            ("UAT acceptance", "30%", "144,000", "On UAT acceptance, target 15 January 2027"),
            ("Go-live", "20%", "96,000", "On go-live, target 31 March 2027")]:
    c = t.add_row().cells
    for i, x in enumerate(row): c[i].text = x
doc.add_paragraph()
h("EXECUTION")
para("Signed for and on behalf of LIM SEONG HAI LIGHTING SDN BHD: Tan Wei Ming, Executive Director — 29 September 2026")
para("Signed for and on behalf of AWAN DIGITAL SOLUTIONS SDN BHD: Farid Osman, Director — 29 September 2026")
doc.save(OUT); print(OUT)
