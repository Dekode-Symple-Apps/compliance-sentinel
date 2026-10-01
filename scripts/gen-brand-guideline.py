"""Renders the Sarawak brand guide (src/lib/brand-guideline.ts) as a Word
document in plain words, so what agencies read and what the AI checks are the
same rules. Official rules cite their circular; the rest are marked good practice.

  python3 scripts/gen-brand-guideline.py  → public/templates/brand/UKAS-BG-01-...docx
"""
import json, os, subprocess
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = subprocess.run(
    ["npx", "tsx", "-e", "import * as g from './src/lib/brand-guideline.ts'; console.log(JSON.stringify({ G: g.GUIDELINE, P: g.PALETTE, T: g.TYPEFACES, C: g.CATEGORIES, L: g.CATEGORY_LABEL, R: g.BRAND_RULES, S: g.SOURCES, F: g.OFFICIAL_FACTS }))"],
    cwd=ROOT, capture_output=True, text=True, check=True).stdout
d = json.loads(src.strip().splitlines()[-1])
G, P, T, C, L, R, S, F = d["G"], d["P"], d["T"], d["C"], d["L"], d["R"], d["S"], d["F"]
RED, BLACK, GREY = RGBColor(0xCE, 0x11, 0x26), RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0x6B, 0x6B, 0x6B)
IF_MISSED = {"critical": "Must fix", "major": "Should fix", "minor": "Tip"}

doc = Document()
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Cm(2.2)
style = doc.styles["Normal"]; style.font.name = "Arial"; style.font.size = Pt(10.5)

def heading(text, size=14, color=BLACK, space=10):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = color
    p.paragraph_format.space_before = Pt(space); return p
def note(text):
    p = doc.add_paragraph(); r = p.add_run(text); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY; return p
def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers): t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = v
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    return t

official = [r for r in R if r["source"]["kind"] == "official"]
t = doc.add_paragraph(); r = t.add_run("KERAJAAN SARAWAK · SARAWAK GOVERNMENT"); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RED
heading(G["title"], 20, BLACK, 4)
doc.add_paragraph(f'{G["version"]} · {G["owner"]}')
note("Draft. No complete Sarawak brand guide is published yet. Rules marked Official come from the State Secretary's circulars; the rest are good practice until UKAS publishes a full guide.")

heading("1. What this guide is")
doc.add_paragraph(f"Every design an agency sends out, from slides and brochures to posters and social posts, is checked against the {len(R)} rules in this guide. {len(official)} rules are official. The check takes about a minute and shows what to fix and where.")

heading("2. What the result means")
table(["Result", "What it means"], [
    ["Ready to publish", "Nothing important is missing. It can go to the brand officer for approval."],
    ["Needs small fixes", "A few things should change. The brand officer may still approve it."],
    ["Must fix before publishing", "Something important is wrong, such as the crest or an official title. Fix it before it goes out."],
], [5, 11])

heading("3. Colours")
table(["Colour", "Use", "Code"], [[c["name"], "Main colour" if c["role"] == "primary" else "Support", c["hex"]] for c in P], [5, 5, 6])
note("No official colour codes are published. These come from the state flag.")
heading("4. Fonts")
doc.add_paragraph(f'Headings: {T["headings"]}. Text: {T["body"]}. If you don\'t have them: {T["fallback"]}.')
note("No official fonts are published. These are placeholders.")

heading("5. Good to know")
for f in F:
    doc.add_paragraph(f'{f["text"]} ({f["ref"]})', style="List Bullet")

heading("6. The rules")
doc.add_paragraph("Must fix: the design should not go out until it is fixed. Should fix: change it unless there is a good reason. Tip: worth doing.")
for cat in C:
    heading(L.get(cat, cat), 12, RED, 8)
    table(["Rule", "What to do", "If missed", "Source"], [
        [r["title"], r["plain"], IF_MISSED[r["severity"]], (f'Official: {r["source"]["ref"]}' if r["source"]["kind"] == "official" else "Good practice")]
        for r in R if r["category"] == cat], [3.6, 7.4, 2.2, 3.6])

heading("7. How approval works")
for line in [
    "The agency uploads the design. It is checked against these rules straight away.",
    "The check shows each thing to fix, with the page and a mark on the spot.",
    "A UKAS brand officer approves it for public use and gives an approval number, or sends it back with what to change.",
    "The agency uploads a new version, which is checked again. Every version and decision is kept.",
]:
    doc.add_paragraph(line, style="List Bullet")

heading("8. Where the official rules come from")
for s in S:
    p = doc.add_paragraph(style="List Bullet"); r = p.add_run(s["ref"]); r.bold = True
    p.add_run(f' — {s["title"]}. {s["issuer"]}, {s["date"]}. {s["url"]}')
note("Circulars are public on eCircular, the Sarawak Government's circular portal (ecircular.sarawak.gov.my).")

heading("Appendix: exact wording for brand officers", 12, BLACK, 14)
doc.add_paragraph("The guideline wording and what the check looks for, by rule number.")
table(["No.", "Rule", "What the check looks for"], [[r["id"], r["rule"], r["check_how"]] for r in R], [1.6, 7.2, 8])

out = os.path.join(ROOT, "public", "templates", "brand")
os.makedirs(out, exist_ok=True)
path = os.path.join(out, f'{G["code"]}-Sarawak-Government-Brand-Guide-{G["version"].replace(" ", "-")}.docx')
doc.save(path)
print(path)
