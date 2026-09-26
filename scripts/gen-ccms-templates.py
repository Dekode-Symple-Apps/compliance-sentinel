"""Build the Commercial CMS contract templates as Word documents.

    python3 scripts/gen-ccms-templates.py

Source of truth is the JSON under src/lib/ccms-templates/ — the same file the
app's deviation check reads — so the document a user downloads and the clause
list a draft is compared against can never drift apart. Output goes to
public/templates/ccms/, served by the app at /templates/ccms/<file>.
"""
import glob, json, os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = f"{PROJ}/src/lib/ccms-templates"
OUT = f"{PROJ}/public/templates/ccms"
GREY = RGBColor(0x59, 0x59, 0x59)


def field(run, instr):
    """Insert a Word field (e.g. PAGE) into a run."""
    for tag, text in (("begin", None), ("instr", instr), ("end", None)):
        if tag == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        run._r.append(el)


def base_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = "Arial"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    pf = st.paragraph_format
    pf.space_after = Pt(6); pf.line_spacing = 1.15
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.4)


def header_footer(doc, t):
    s = doc.sections[0]
    hp = s.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("PRIVATE & CONFIDENTIAL"); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = GREY
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"{t['code']} · {t['title']} · v{t['version']}   |   Page ")
    r.font.size = Pt(8); r.font.color.rgb = GREY
    r2 = fp.add_run(); r2.font.size = Pt(8); r2.font.color.rgb = GREY; field(r2, "PAGE")


def para(doc, text, bold=False, size=None, align=None, space_after=None, italic=False, indent=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align is not None: p.alignment = align
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    if indent is not None: p.paragraph_format.left_indent = Cm(indent)
    return p


def usage_page(doc, t):
    para(doc, "TEMPLATE NOTE — DELETE THIS PAGE BEFORE ISSUE", bold=True, size=9)
    para(doc, f"{t['code']}  {t['title']}  (version {t['version']}, {t['effectiveDate']})", size=9)
    para(doc, f"Owner: {t['owner']}   ·   Status: {t['status']}", size=9)
    para(doc, t["usageNote"], size=9)
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for c, h in zip(hdr, ("Clause", "Title", "Status", "Approved position")):
        c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(8)
    for cl in t["clauses"]:
        row = tbl.add_row().cells
        status = "LOCKED" if cl["locked"] else ("Mandatory" if cl["mandatory"] else "Optional")
        for c, v in zip(row, (cl["number"], cl["title"], status, cl["keyPosition"])):
            c.text = ""; r = c.paragraphs[0].add_run(v); r.font.size = Pt(8)
            if v == "LOCKED": r.bold = True
    widths = (Cm(1.4), Cm(4.2), Cm(2.0), Cm(9.0))
    tbl.autofit = False
    for col, w in zip(tbl.columns, widths):
        col.width = w
    for row in tbl.rows:
        for c, w in zip(row.cells, widths):
            c.width = w
    para(doc, "")
    para(doc, "Assumptions to confirm before adoption:", bold=True, size=9)
    for a in t["assumptions"]:
        para(doc, f"•  {a}", size=9, indent=0.4, space_after=2)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def body(doc, t):
    para(doc, t["title"].upper(), bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "(Confidentiality Agreement)", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    para(doc, "THIS AGREEMENT is made on the date stated in Schedule 1", space_after=10)
    para(doc, "BETWEEN", bold=True)
    para(doc, "(1)  " + t["parties"]["company"] + "; and", indent=0.5)
    para(doc, "AND", bold=True)
    para(doc, "(2)  " + t["parties"]["counterparty"] + ",", indent=0.5)
    para(doc, "(each a \"Party\" and together the \"Parties\").", space_after=10)
    para(doc, "RECITALS", bold=True)
    for i, rc in enumerate(t["recitals"]):
        para(doc, f"({chr(65 + i)})  {rc}", indent=0.5)
    para(doc, "IT IS AGREED as follows:", bold=True, space_after=10)
    for cl in t["clauses"]:
        h = para(doc, f"{cl['number']}.  {cl['title'].upper()}", bold=True, space_after=4)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.space_before = Pt(8)
        for ptxt in cl["paragraphs"]:
            # Definitions (quoted term first) sit a step further in under 1.1.
            para(doc, ptxt, indent=1.5 if ptxt.startswith('"') else 0.9, align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def schedule(doc, t):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    sc = t["schedule"]
    para(doc, sc["title"].upper(), bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    tbl = doc.add_table(rows=0, cols=3); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for it in sc["items"]:
        row = tbl.add_row().cells
        for c, v, b in zip(row, (it["number"], it["label"], it["value"]), (False, True, False)):
            c.text = ""; r = c.paragraphs[0].add_run(v); r.bold = b; r.font.size = Pt(9.5)
    tbl.autofit = False
    for col, w in zip(tbl.columns, (Cm(1.0), Cm(5.0), Cm(10.6))):
        col.width = w
    for row in tbl.rows:
        for c, w in zip(row.cells, (Cm(1.0), Cm(5.0), Cm(10.6))):
            c.width = w


def execution(doc, t):
    ex = t["execution"]
    para(doc, "")
    para(doc, ex["intro"], space_after=12)
    tbl = doc.add_table(rows=1, cols=2)
    for cell, blk in zip(tbl.rows[0].cells, ex["blocks"]):
        cell.text = ""
        r = cell.paragraphs[0].add_run(blk["party"]); r.bold = True; r.font.size = Pt(9.5)
        cell.add_paragraph()
        cell.add_paragraph("_______________________________")
        for ln in blk["lines"]:
            p = cell.add_paragraph(); rr = p.add_run(ln); rr.font.size = Pt(9)
            if ln.startswith("In the presence"):
                cell.add_paragraph(); cell.add_paragraph("_______________________________")


def build(path):
    t = json.load(open(path, encoding="utf-8"))
    doc = Document()
    base_styles(doc)
    header_footer(doc, t)
    usage_page(doc, t)
    body(doc, t)
    schedule(doc, t)
    execution(doc, t)
    os.makedirs(OUT, exist_ok=True)
    name = f"{t['code']}-{t['title'].replace(' ', '-')}-v{t['version']}.docx"
    doc.save(f"{OUT}/{name}")
    return name


if __name__ == "__main__":
    for p in sorted(glob.glob(f"{SRC}/*.json")):
        print("wrote", build(p))
