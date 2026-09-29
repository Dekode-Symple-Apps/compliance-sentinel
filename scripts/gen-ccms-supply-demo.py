"""The straight-through demo: a routine supply agreement with an approved,
low-risk vendor, on complete and balanced terms — the AI finds it clean, so it
is cleared and approved without waiting, then signed, stamped and filed.

  python3 scripts/gen-ccms-supply-demo.py →
    ~/Desktop/01. Demo Data/Commercial CMS/2 Straight-through - Sinar Precast supply/
      Step 1 - Upload Draft on New Request - supply agreement.docx
      Step 2 - Upload Signed Copy - signed and stamped.docx
Invented details; entity particulars are the demo values in src/lib/ccms.ts.
"""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

D = os.path.expanduser("~/Desktop/01. Demo Data/Commercial CMS/2 Straight-through - Sinar Precast supply/")
os.makedirs(D, exist_ok=True)

CLAUSES = [
    ("Supply", [
        "The Supplier shall supply the precast concrete U-drains, covers and sumps described in Schedule 1 (the \"Goods\") to the LSH 33 Block C site, Jalan Klang Lama, Kuala Lumpur.",
        "The Goods shall comply with MS 1293, be SIRIM certified and carry a valid CIDB Perakuan Pematuhan Standard (PPS) under section 33D of the Construction Industry Development Board Act 1994. Each delivery shall be accompanied by the manufacturer's test certificates and a copy of the PPS.",
    ]),
    ("Price", [
        "The contract price is RM186,500 (Ringgit Malaysia One Hundred and Eighty-Six Thousand Five Hundred), fixed for the duration of this Agreement and exclusive of SST.",
        "The price includes delivery to site, unloading and packaging.",
    ]),
    ("Delivery", [
        "The Supplier shall deliver the Goods in the batches and on the dates in Schedule 2, the last batch by 30 November 2026.",
        "The Purchaser may suspend or reschedule a batch by giving the Supplier at least 7 days' written notice, without additional cost if the new date is within 60 days of the original date.",
        "If the Supplier fails to deliver a batch on time, it shall pay liquidated damages of 0.1% of the value of the late batch for each day of delay, capped at 10% of the contract price.",
    ]),
    ("Inspection and Acceptance", [
        "The Purchaser may inspect the Goods on delivery and reject non-conforming Goods by written notice within 7 days. The Supplier shall replace rejected Goods within 14 days at its own cost.",
        "If the Supplier fails to deliver or replace Goods when required, the Purchaser may, after 7 days' written notice, buy substitute goods from a third party and recover any additional cost from the Supplier.",
    ]),
    ("Payment", [
        "The Supplier shall invoice each batch after delivery and acceptance. The Purchaser shall pay each valid invoice within 30 days of receipt.",
        "The Purchaser may set off against any payment due to the Supplier any liquidated damages, substitute-goods costs or other amounts the Supplier owes under this Agreement.",
    ]),
    ("Title and Risk", ["Title and risk in the Goods pass to the Purchaser on acceptance at site."]),
    ("Warranty", ["The Supplier warrants the Goods against defects in materials and workmanship for 12 months from acceptance and shall repair or replace defective Goods at its own cost."]),
    ("Insurance", ["The Supplier shall maintain goods-in-transit insurance and product liability insurance of at least RM1,000,000 per occurrence for the duration of this Agreement and the warranty period, and shall provide certificates of insurance on request."]),
    ("Indemnity and Liability", [
        "The Supplier shall indemnify the Purchaser against third-party claims arising from defects in the Goods or from the Supplier's negligence or breach of this Agreement.",
        "Each party's total liability under this Agreement is limited to the contract price, except that no limit applies to death or personal injury, damage to property, fraud, gross negligence, wilful misconduct, the indemnity in this clause, or breach of the Anti-Bribery and Corruption clause.",
    ]),
    ("Anti-Bribery and Corruption", [
        "The Supplier shall comply with the Malaysian Anti-Corruption Commission Act 2009 (including section 17A) and the Purchaser's Anti-Bribery and Corruption Policy, and shall not offer or give any gratification in connection with this Agreement.",
        "The Purchaser may terminate this Agreement immediately by written notice if the Supplier breaches this clause.",
    ]),
    ("Confidentiality and Personal Data", ["Each party shall keep the other's confidential information confidential and shall comply with the Personal Data Protection Act 2010."]),
    ("Termination", [
        "Either party may terminate this Agreement if the other commits a material breach and does not remedy it within 14 days of written notice.",
        "The Purchaser may terminate this Agreement immediately by written notice if the Supplier becomes insolvent, enters into administration, liquidation or winding-up, or makes an arrangement with its creditors.",
        "The Purchaser may terminate this Agreement for convenience on 14 days' written notice, paying for Goods delivered and accepted and the reasonable cost of Goods already manufactured for this Agreement.",
    ]),
    ("Force Majeure", ["Neither party is liable for delay caused by events beyond its reasonable control, provided it notifies the other promptly and uses reasonable efforts to mitigate. If such an event continues for more than 30 days, either party may terminate by written notice."]),
    ("Assignment and Subcontracting", ["The Supplier shall not assign this Agreement or subcontract the manufacture of the Goods without the Purchaser's prior written consent."]),
    ("Notices", [
        "Notices shall be in writing and delivered by hand, registered post or email to: the Purchaser — attention Head of Procurement, Wisma Lim Seong Hai, Kuala Lumpur, procurement@lsh-best.example; the Supplier — attention Rahman bin Ali, General Manager, Lot 45, Kawasan Perindustrian Senai, Johor, rahman@sinarprecast.example.",
    ]),
    ("Stamp Duty", ["The Supplier shall bear the stamp duty payable on this Agreement under the Stamp Act 1949 and arrange stamping within 30 days of execution."]),
    ("Governing Law and Disputes", ["This Agreement is governed by the laws of Malaysia and the parties submit to the exclusive jurisdiction of the courts of Malaysia."]),
    ("Entire Agreement", ["This Agreement, with its Schedules, is the entire agreement between the parties on its subject matter and may be varied only in writing signed by both parties."]),
]


def build(path, signed):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    def h(t, size=12, center=False):
        p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(size)
        if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para = lambda t: doc.add_paragraph(t)

    h("SUPPLY AGREEMENT", 16, True)
    h("Precast Concrete Drains — LSH 33 Block C", 12, True)
    if signed: para("Stamp duty paid — Certificate No. STAMPS-2026-DEMO-0231 · LHDN · 28 September 2026")
    para("THIS AGREEMENT is made on 28 September 2026 and takes effect on 1 October 2026 (the \"Commencement Date\").")
    h("BETWEEN")
    para("(1) LSH BEST BUILDERS SDN BHD (Registration No. 200501023456 (234567-A)), of Wisma Lim Seong Hai, Kuala Lumpur (the \"Purchaser\"); and")
    para("(2) SINAR PRECAST INDUSTRIES SDN BHD (Registration No. 201501012345), of Lot 45, Kawasan Perindustrian Senai, Johor (the \"Supplier\").")
    for i, (title, lines) in enumerate(CLAUSES, 1):
        h(f"{i}. {title}")
        for j, l in enumerate(lines, 1): para(f"{i}.{j} {l}")
    h("SCHEDULE 1 — GOODS")
    t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
    for i, x in enumerate(["Item", "Quantity", "Unit price (RM)", "Amount (RM)"]): t.rows[0].cells[i].text = x
    for row in [("U-drain 600 mm x 600 mm x 1.2 m", "620", "185", "114,700"), ("Drain cover, heavy duty", "620", "95", "58,900"), ("Sump 1.2 m x 1.2 m", "30", "430", "12,900")]:
        c = t.add_row().cells
        for i, x in enumerate(row): c[i].text = x
    c = t.add_row().cells; c[0].text = "Total"; c[3].text = "186,500"
    h("SCHEDULE 2 — DELIVERY")
    para("Batch 1 (40%): by 23 October 2026 · Batch 2 (40%): by 13 November 2026 · Batch 3 (20%): by 30 November 2026")
    h("EXECUTION")
    para("Executed in accordance with section 66 of the Companies Act 2016.")
    for party, (a, b) in [("LSH BEST BUILDERS SDN BHD", (("Tan Wei Ming", "Director"), ("Lee Mei Fong", "Director"))),
                          ("SINAR PRECAST INDUSTRIES SDN BHD", (("Rahman bin Ali", "Director"), ("Siti Aminah binti Yusof", "Company Secretary")))]:
        para(f"Signed for and on behalf of {party}:")
        if signed:
            para(f"    {a[0]}, {a[1]} — 28 September 2026        {b[0]}, {b[1]} — 28 September 2026")
        else:
            para(f"    ____________________ Name / {a[1]} / Date        ____________________ Name / {b[1]} / Date")
    doc.save(path); print(path)


build(D + "Step 1 - Upload Draft on New Request - supply agreement.docx", signed=False)
build(D + "Step 2 - Upload Signed Copy - signed and stamped.docx", signed=True)
